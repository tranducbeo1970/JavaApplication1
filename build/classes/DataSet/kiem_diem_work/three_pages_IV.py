from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from copy import deepcopy
root=Path(r'C:\Users\tranduc\Downloads\duc')
d=Document(next(root.glob('2.*')))
top=deepcopy(d.tables[0]._tbl);sig=deepcopy(d.tables[1]._tbl)
body=d._element.body
for c in list(body):
 if not c.tag.endswith('}sectPr'):body.remove(c)
body.insert(0,top)
def p(text,bold=False,center=False):
 q=d.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
 f=q.paragraph_format;f.space_after=Pt(4);f.space_before=Pt(0);f.line_spacing=Pt(16);f.widow_control=True
 f.keep_with_next=bold and not center
 r=q.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(13);r.bold=bold
 return q
p('BẢN KIỂM ĐIỂM CÁ NHÂN',True,True)
p('Về trách nhiệm liên quan đến Kết luận thanh tra số 430/KL-TTCP ngày 16/7/2026 của Thanh tra Chính phủ',False,True)
p('Họ và tên: Trần Đức')
p('Chức vụ về Đảng: ........................................................')
p('Chức vụ chính quyền: Trưởng Phòng Nghiên cứu phát triển theo thời kỳ')
p('Chức vụ liên quan đến nội dung kiểm điểm: Trưởng Phòng Nghiên cứu phát triển trong thời kỳ có liên quan.')
p('I. CHỨC TRÁCH, NHIỆM VỤ VÀ PHẠM VI KIỂM ĐIỂM',True)
p('Trong thời gian giữ chức vụ Trưởng phòng Nghiên cứu phát triển (NCPT), tôi có trách nhiệm quản lý, điều hành Phòng; kiểm tra tiến độ, chất lượng và hồ sơ nhiệm vụ khoa học và công nghệ (KH&CN) do Phòng chủ trì; tham gia lập và tổ chức thực hiện kế hoạch được giao.')
p('Đối chiếu Phụ lục I Kế hoạch số 383, trách nhiệm của cá nhân tôi được xem xét tập trung đối với 02 nội dung:')
p('- Mã I-5.1c - Quản lý và thực hiện nhiệm vụ biển báo LED nền dẫn sáng, với trách nhiệm người đứng đầu Phòng NCPT là cơ quan chủ trì.')
p('- Mã I-5.1a - Tham gia lập Kế hoạch hoạt động KH&CN và tổ chức triển khai các nhiệm vụ được giao cho Phòng NCPT.')
p('Hai nội dung cùng đối chiếu Phụ lục số 02-II-3.1 của Kết luận số 430/KL-TTCP, trong phạm vi công việc và thời kỳ tôi phụ trách.')
p('II. KIỂM ĐIỂM THEO TỪNG NỘI DUNG',True)
p('1. Quản lý cơ quan chủ trì nhiệm vụ KH&CN (Mã I-5.1c)',True)
p('1.1. Nội dung và hồ sơ có liên quan',True)
p('Phòng NCPT chủ trì nhiệm vụ “Nghiên cứu thiết kế chế tạo biển báo công nghệ mới sử dụng LED nền dẫn sáng”. Quá trình triển khai còn tồn tại về tiến độ, tính đầy đủ của hồ sơ; một số nội dung phải chỉnh sửa, bổ sung nhiều lần, chưa được xử lý dứt điểm kịp thời.')
d.add_page_break()
p('1.2. Vai trò và thiếu sót của cá nhân',True)
p('Tôi còn dựa nhiều vào cán bộ được phân công, chưa thường xuyên kiểm tra đến cùng chất lượng và tính đầy đủ của hồ sơ. Việc nắm bắt, đôn đốc xử lý khó khăn có lúc chưa kịp thời, chưa quyết liệt; kiểm tra sau phân công chưa chặt chẽ. Tôi nhận đây là thiếu sót chủ quan trong thực hiện trách nhiệm người đứng đầu.')
p('1.3. Thông tin và khả năng phát hiện, ngăn ngừa',True)
p('Qua hồ sơ và báo cáo tiến độ, nếu chủ động đặt các mốc kiểm soát, yêu cầu báo cáo từng đầu việc và trực tiếp kiểm tra nội dung quan trọng, tôi có thể phát hiện, chấn chỉnh sớm hơn một số tồn tại.')
p('1.4. Nguyên nhân, ảnh hưởng và tự xác định trách nhiệm',True)
p('Bản thân chưa sâu sát, chưa chủ động nhận diện nguy cơ chậm tiến độ và sai sót hồ sơ, ảnh hưởng đến chất lượng quản lý nhiệm vụ. Tôi nhận trách nhiệm quản lý đối với các tồn tại thuộc Phòng NCPT trong thời kỳ phụ trách.')
p('2. Lập và triển khai Kế hoạch hoạt động KH&CN (Mã I-5.1a)',True)
p('2.1. Nội dung và hồ sơ có liên quan',True)
p('Phòng NCPT tham gia đề xuất kế hoạch và thực hiện nhiệm vụ được giao. Qua rà soát, kế hoạch một số năm phải điều chỉnh; một số nhiệm vụ chưa hoàn thành theo thời gian dự kiến.')
p('2.2. Vai trò và thiếu sót của cá nhân',True)
p('Tôi chưa chỉ đạo đánh giá đầy đủ tính khả thi, nguồn lực và thời gian trước khi đề xuất. Việc bố trí nguồn lực, kiểm tra, đôn đốc một số công việc chưa sát yêu cầu; xử lý nguy cơ chậm tiến độ có lúc chưa kịp thời.')
p('2.3. Thông tin và khả năng phát hiện, ngăn ngừa',True)
p('Từ khối lượng công việc và nguồn lực của Phòng, bản thân có thể rà soát kỹ hơn, dự báo khó khăn để đề xuất kế hoạch sát thực tế và chủ động giải pháp khi triển khai.')
p('2.4. Nguyên nhân, ảnh hưởng và tự xác định trách nhiệm',True)
p('Việc dự báo và kiểm soát tiến độ của tôi còn hạn chế, ảnh hưởng đến chất lượng đề xuất và tổ chức thực hiện. Tôi nhận trách nhiệm người đứng đầu đối với phần việc thuộc đơn vị trong thời kỳ phụ trách.')
d.add_page_break()
p('III. NGUYÊN NHÂN CHUNG VÀ BÀI HỌC TRÁCH NHIỆM',True)
p('1. Khách quan: Nhiệm vụ nghiên cứu có yêu cầu kỹ thuật phức tạp, nhân sự có biến động và kiêm nhiệm nhiều công việc. Tôi nhận thức các yếu tố này không loại trừ trách nhiệm quản lý của mình.')
p('2. Chủ quan: Bản thân chưa sâu sát, chưa quyết liệt; chưa gắn phân công với kiểm tra kết quả; chưa coi trọng đầy đủ việc kiểm soát hồ sơ song song với chuyên môn. Tôi nghiêm túc nhận thức người đứng đầu phải chủ động phát hiện, chấn chỉnh thiếu sót và chịu trách nhiệm về kết quả công việc đã phân công.')
p('IV. BIỆN PHÁP KHẮC PHỤC',True)
p('Mặc dù hiện không còn giữ chức vụ Trưởng phòng NCPT, tôi nhận thức rõ trách nhiệm tiếp tục góp phần khắc phục thiếu sót trong thời kỳ phụ trách và nghiêm túc thực hiện:')
p('Một là, chủ động tham mưu lãnh đạo biện pháp khắc phục tồn tại; phối hợp với Phòng NCPT và các đơn vị liên quan rà soát, bổ sung hồ sơ, cung cấp thông tin và thực hiện các phần việc được giao.')
p('Hai là, vận dụng bài học kiểm điểm vào chức năng, nhiệm vụ mới; nâng cao chất lượng tham mưu, làm việc sâu sát, kiểm tra kỹ hồ sơ và theo dõi tiến độ công việc của bản thân.')
p('Ba là, kịp thời báo cáo, đề xuất xử lý vướng mắc; nghiêm túc tiếp thu góp ý, chủ động khắc phục hạn chế, không để lặp lại thiếu sót trước đây trong nhiệm vụ mới.')
p('V. TỰ NHẬN TRÁCH NHIỆM VÀ ĐỀ XUẤT XỬ LÝ',True)
p('Qua kiểm điểm, tôi nhận thức rõ trách nhiệm người đứng đầu về chất lượng, tiến độ công việc và tính đầy đủ của hồ sơ thuộc đơn vị.')
p('Tôi nghiêm túc nhận thiếu sót trong quản lý, kiểm tra, đôn đốc; chưa chủ động phòng ngừa và xử lý kịp thời một số tồn tại trong thời kỳ phụ trách.')
p('Đối với mã I-5.1c, tôi tự nhận trách nhiệm quản lý cơ quan chủ trì nhiệm vụ LED về việc kiểm tra hồ sơ và đôn đốc xử lý vướng mắc chưa chặt chẽ, chưa kịp thời.')
p('Đối với mã I-5.1a, tôi tự nhận trách nhiệm về việc chỉ đạo đánh giá nguồn lực, tính khả thi khi đề xuất kế hoạch chưa sát thực tế và tổ chức thực hiện một số nhiệm vụ chưa đáp ứng tiến độ.')
p('Tôi nghiêm túc tiếp thu góp ý, cam kết khắc phục và không để lặp lại hạn chế tương tự. Kính đề nghị cấp có thẩm quyền xem xét trách nhiệm, hình thức xử lý trên cơ sở tính chất, mức độ thiếu sót, hồ sơ thực tế và thời kỳ phụ trách. Tôi nghiêm túc chấp hành kết luận.')
p('Tôi xin nghiêm túc kiểm điểm.')
q=p('NGƯỜI KIỂM ĐIỂM',True);q.alignment=WD_ALIGN_PARAGRAPH.RIGHT
q=p('Trần Đức',True);q.alignment=WD_ALIGN_PARAGRAPH.RIGHT;q.paragraph_format.space_before=Pt(24)
out=Path(__file__).parent.parent/'Ban_kiem_diem_Tran_Duc_3_trang_sua_muc_IV.docx';d.save(out);print(out)
