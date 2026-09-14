from pathlib import Path
from docx import Document
import ast
source=Document(r'C:\Users\tranduc\Downloads\duc\KD_V2.docx')
texts=[p.text for p in source.paragraphs]
date_path=Path(__file__).with_name('plan_383_date.txt')
plan='383'
content=[
(2,'BẢN KIỂM ĐIỂM CÁ NHÂN'),
(3,'Về trách nhiệm của ông Trần Đức liên quan đến nội dung trong Kết luận thanh tra số 430/KL-TTCP ngày 16/7/2026 của Thanh tra Chính phủ'),
(6,'Họ và tên: Trần Đức'),
(7,'Chức vụ về Đảng: ........................................................'),
(8,'Chức vụ chính quyền: Trưởng Phòng Nghiên cứu phát triển theo thời kỳ'),
(9,'Chức vụ liên quan đến nội dung kiểm điểm: Trưởng Phòng Nghiên cứu phát triển trong thời kỳ có liên quan; người đứng đầu cơ quan chủ trì nhiệm vụ KH&CN, đơn vị tham gia lập và triển khai Kế hoạch hoạt động KH&CN.'),
]
changes={
5:f'Đối chiếu Phụ lục I Kế hoạch số {plan}, trách nhiệm của cá nhân tôi được xem xét tập trung đối với 02 nội dung:',
6:'- Mã I-5.1c - Quản lý và thực hiện nhiệm vụ KH&CN “Nghiên cứu thiết kế chế tạo biển báo công nghệ mới sử dụng LED nền dẫn sáng”: trách nhiệm quản lý của người đứng đầu Phòng NCPT là cơ quan chủ trì nhiệm vụ.',
7:'- Mã I-5.1a - Lập và triển khai Kế hoạch hoạt động KH&CN: trách nhiệm của Trưởng Phòng NCPT trong việc tham gia lập kế hoạch và tổ chức triển khai các nhiệm vụ được giao cho Phòng.',
10:'1. Nội dung quản lý cơ quan chủ trì nhiệm vụ KH&CN (MÃ I-5.1c)',
33:'2. Nội dung lập và triển khai Kế hoạch hoạt động KH&CN (MÃ I-5.1a)',
71:'Qua kiểm điểm, tôi nhận thức rõ trách nhiệm của người đứng đầu Phòng NCPT đối với chất lượng, tiến độ và tính đầy đủ của hồ sơ nhiệm vụ KH&CN; trách nhiệm trong việc tham gia lập kế hoạch sát thực tế và tổ chức triển khai nhiệm vụ được giao. Tôi nghiêm túc nhìn nhận những hạn chế của bản thân trong thực hiện chức trách, nhiệm vụ.',
72:'Tôi nghiêm túc nhận thiếu sót về việc có thời điểm chưa sâu sát, chưa quyết liệt trong quản lý, kiểm tra và đôn đốc; chưa gắn việc phân công với kiểm soát kết quả đầy đủ; chưa kịp thời phát hiện, phòng ngừa và xử lý các tồn tại về hồ sơ, tiến độ thuộc phạm vi Phòng NCPT trong thời kỳ tôi phụ trách.',
73:'Đối với mã I-5.1c, tôi tự nhận trách nhiệm quản lý của người đứng đầu Phòng NCPT là cơ quan chủ trì nhiệm vụ biển báo LED về việc kiểm tra, giám sát tính đầy đủ của hồ sơ và đôn đốc xử lý vướng mắc có thời điểm chưa chặt chẽ, chưa kịp thời. Tôi nhận thấy bản thân chưa thực hiện đầy đủ vai trò kiểm tra sau phân công để hạn chế các thiếu sót phát sinh trong phạm vi đơn vị và thời kỳ phụ trách.',
74:'Đối với mã I-5.1a, tôi tự nhận trách nhiệm của Trưởng Phòng NCPT về việc chỉ đạo đánh giá tính khả thi, nguồn lực và tiến độ khi tham gia lập Kế hoạch hoạt động KH&CN có nội dung chưa sát thực tế; công tác tổ chức, kiểm tra và đôn đốc thực hiện một số nhiệm vụ được giao cho Phòng còn hạn chế, chưa đáp ứng tiến độ dự kiến.',
75:'Căn cứ tính chất, mức độ, vai trò và hậu quả hiện được xác định, tôi tự nhận đây là những thiếu sót trong công tác quản lý, kiểm tra, đôn đốc và tổ chức thực hiện nhiệm vụ thuộc trách nhiệm của bản thân. Tôi nghiêm túc tiếp thu, rút kinh nghiệm và cam kết khắc phục, không để lặp lại những hạn chế tương tự.',
77:'Kính đề nghị cấp có thẩm quyền xem xét mức độ trách nhiệm, hình thức xử lý (nếu có) trên cơ sở kết quả kiểm điểm của đầy đủ các Nhóm, hồ sơ chứng minh và mối quan hệ nhân quả của từng hành vi, gắn với chức trách, nhiệm vụ và thời kỳ phụ trách của cá nhân tôi. Tôi nghiêm túc chấp hành kết luận của cấp có thẩm quyền.'
}
for i in range(3,79):
 t=changes.get(i,texts[i])
 if i in [3,9,52,62,70]:role={3:10,9:15,52:26,62:31,70:37}[i]
 elif i in [10,33,53,56] or (len(t)>3 and t[0].isdigit() and t[1]=='.'):role=17
 elif i in [6,7]:role=13
 else:role=18
 content.append((role,t))
 if i==7:content.append((18,'Hai mã nêu trên thuộc Phụ lục I của Kế hoạch số 383, cùng đối chiếu với nội dung tại Phụ lục số 02-II-3.1 của Kết luận thanh tra số 430/KL-TTCP.'))
from zipfile import ZipFile
from lxml import etree as E
from copy import deepcopy
import hashlib
ref=next(Path(r'C:\Users\tranduc\Downloads\duc').glob('2.*'))
z=ZipFile(ref)
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W='{'+ns['w']+'}'
tree=E.fromstring(z.read('word/document.xml'));body=tree.find('w:body',ns)
ps=body.findall('w:p',ns); tables=body.findall('w:tbl',ns)
top=deepcopy(tables[0]); sig=deepcopy(tables[1]);sect=deepcopy(body.find('w:sectPr',ns))
def para(idx,text):
 p=deepcopy(ps[idx]);rp=p.find('w:r/w:rPr',ns)
 for c in list(p):
  if c.tag!=W+'pPr':p.remove(c)
 r=E.SubElement(p,W+'r')
 if rp is not None:r.append(deepcopy(rp))
 E.SubElement(r,W+'t').text=text
 if idx in [10,15,17,26,27,29,31,37]:
  pp=p.find('w:pPr',ns)
  if pp is None:pp=E.SubElement(p,W+'pPr')
  if pp.find('w:keepNext',ns) is None:E.SubElement(pp,W+'keepNext')
 return p
for c in list(body):body.remove(c)
body.append(top)
for idx,t in content:body.append(para(idx,t))
last=sig.findall('w:tr/w:tc',ns)[1].findall('w:p',ns)[-1]
r=E.SubElement(last,W+'r');rp=E.SubElement(r,W+'rPr');E.SubElement(rp,W+'b');E.SubElement(r,W+'t').text='Trần Đức'
body.append(sig);body.append(sect)
final=Path(__file__).parent.parent/'Ban_kiem_diem_Tran_Duc_KD_V2_dung_mau_va_ma.docx'
with ZipFile(final,'w') as zz:
 for info in z.infolist():zz.writestr(info,E.tostring(tree,encoding='UTF-8',xml_declaration=True,standalone=True) if info.filename=='word/document.xml' else z.read(info.filename))
Path(__file__).with_name('artifact_correct_codes.md').write_text('Reference: '+str(ref)+'\nSHA256: '+hashlib.sha256(ref.read_bytes()).hexdigest()+'\nPreserved all package parts except document.xml. Preserved source section geometry, national heading and signature tables. Body cloned from template paragraph roles with KD_V2 text; revised I, II codes and V per user; plan date omitted at user direction. Party position and signature date remain unconfirmed. Rendering unavailable because LibreOffice missing.\n',encoding='utf8')
final=Path(__file__).parent.parent/'Ban_kiem_diem_Tran_Duc_KD_V2_dung_mau_va_ma.docx'
d=Document(final)
for p in d.paragraphs:
 if 'Mã ' in p.text or 'Đối với mã' in p.text or 'Đối chiếu' in p.text:print(p.text)
assert all(sum(k in p.text for p in d.paragraphs)>=3 for k in ['I-5.1a','I-5.1c'])
